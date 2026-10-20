"""Build the shared online error registry and administrator Word handbook."""
import ast,hashlib,json,sys
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT/'app'))
import admin_ops

def template(node):
 if isinstance(node,ast.Constant) and isinstance(node.value,str):return node.value
 if isinstance(node,ast.JoinedStr):return ''.join(part.value if isinstance(part,ast.Constant) else '{detail}' for part in node.values)
 return None

def build():
 entries={}
 for path in [ROOT/'app/server.py',ROOT/'app/admin_ops.py']:
  for node in ast.walk(ast.parse(path.read_text(encoding='utf-8'))):
   if not isinstance(node,ast.Call) or len(node.args)<2:continue
   if not ((isinstance(node.func,ast.Name) and node.func.id=='Problem') or (isinstance(node.func,ast.Attribute) and node.func.attr=='Problem')):continue
   status=node.args[0].value if isinstance(node.args[0],ast.Constant) else None;message=template(node.args[1])
   if status not in admin_ops.FAMILIES or not message:continue
   code='PP-'+str(status)+'-'+hashlib.sha256((str(status)+':'+message).encode()).hexdigest()[:6].upper()
   family=admin_ops.FAMILIES[status]
   checks=family[3];meaning=family[2]
   lower=message.lower()
   if 'password' in lower and status==401:checks='Confirm the account name and password privately. Do not put passwords into an error report. Check whether the account is disabled or pending approval.'
   elif 'origin' in lower or 'refresh this page' in lower:checks='Refresh the current HTTPS page. Verify its domain and CSRF session. Check reverse-proxy origin configuration; do not disable these checks.'
   elif 'image' in lower or 'photo' in lower:checks='Check the selected image, its ownership and file limits. For generation failures, run diagnostics and check verified models, memory and dependencies.'
   elif 'permission' in lower or 'administrator' in lower:checks='Check the account role, fine-grained grants and disabled state. A full administrator can correct access. Retain one enabled full administrator.'
   elif 'recipe' in lower or 'export' in lower:checks='Check the selected provider, supported file format, recipe title and ingredient list. Review archive and field size limits. Keep the original export.'
   elif 'section' in lower:checks='Review the Home inventory section name, display order, duplicate names and archived state. Reload current section definitions before retrying.'
   entries[code]={'code':code,'status':status,'template':message,'message':message,'meaning':meaning,'checks':checks,'category':family[1]}
 result=sorted(entries.values(),key=lambda x:x['code'])
 (ROOT/'app/error_catalog.json').write_text(json.dumps(result,indent=2,ensure_ascii=False),encoding='utf-8')
 admin_ops._catalog=None
 return sorted(admin_ops.catalog().values(),key=lambda x:(x['status'],x['code']))

def document(entries):
 from docx import Document
 from docx.shared import Inches,Pt,RGBColor
 from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
 from docx.oxml import OxmlElement
 from docx.oxml.ns import qn
 doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
 sec.top_margin=sec.bottom_margin=Inches(.7);sec.left_margin=sec.right_margin=Inches(.75)
 for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
  st=doc.styles[name];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0)
 doc.styles['Normal'].font.size=Pt(11);doc.styles['Normal'].paragraph_format.space_after=Pt(7)
 doc.styles['Title'].font.size=Pt(24);doc.styles['Heading 1'].font.size=Pt(16)
 for node in doc.styles.element.iter(qn('w:pBdr')):node.getparent().remove(node)
 for name in ['Title','Subtitle','Heading 1','Heading 2']:doc.styles[name].font.underline=False
 header=sec.header.paragraphs[0];header.text='PANTRIPORTAL   ADMINISTRATOR REFERENCE';header.style='Caption';header.runs[0].font.color.rgb=RGBColor(0,0,0)
 footer=sec.footer.paragraphs[0];footer.text='PantriPortal 0.9   |   October 19 2026   |   Page '
 field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
 doc.add_heading('PantriPortal administrator error guide',0)
 doc.add_paragraph('Operations and troubleshooting reference',style='Subtitle')
 doc.add_paragraph('Use this guide to interpret application errors, locate the matching activity event, and choose a practical next check. The online Error reference uses the same registry. Click an error code in Activity and error logs to open its guidance directly.')
 doc.add_heading('Find and investigate a problem',1)
 for text in ['Record the error code and request reference displayed with the failed action. Do not include passwords, session cookies, invite codes or setup tokens in a report.','Open My dashboard then Administration then Activity and error logs. Search the reference, account, action, date range or error code. Errors only narrows the list.','Click the error code for its message, meaning and checks. Run Health diagnostics when the cause is unclear. Use the same reference when inspecting private host service logs.','Correct the cause and retry the smallest affected action. Reload before resolving concurrent edits. Preserve the original file and current database before recovery work.']:doc.add_paragraph(text,style='List Number')
 doc.add_heading('Access and safe recovery',1)
 doc.add_paragraph('Existing administrators keep their abilities. Full administrators can assign platform roles and fine-grained Allow or Deny rules. Household membership is separate. User management cannot disable or demote the last enabled full administrator. Author identifies a recipe contributor; it does not grant admin access by itself.')
 doc.add_paragraph('Database and recovery shows table counts and disk information without exposing record contents or credential hashes. A full administrator can create a verified private recovery archive on the host. Restore requires private host access and a maintenance window. Browser inventories and unsaved drafts are not included in the server backup.')
 doc.add_heading('What the activity log records',1)
 doc.add_paragraph('Logging begins when this release is installed. It cannot reconstruct earlier activity. Events include a timestamp, authenticated actor when known, operation, normalized route, record identifier when available, result status, error code, request reference and duration. Background image failures and signed-in browser script error types are recorded too. Failed sign-ins are labeled as attempts rather than authenticated actions. Unexpected failures include file names, function names and line numbers without variable values or raw exception messages. Passwords, tokens, form contents and URL queries are not logged. Administrator access to logs is permission controlled.')
 doc.add_paragraph('Audit records are append only through the application and signed as a sequence with a private host key. Verify audit integrity checks that sequence. This detects altered rows; it does not protect against an attacker who controls both the database and signing key. Preserve the key with private backups. There is no automatic log deletion in this release. Watch disk space and export filtered CSV batches of up to 5000 events. Private host administrators must plan retention and secure archives.')
 doc.add_heading('Recipe connections and Home inventory',1)
 doc.add_paragraph('Recipe connections supports the vetted public-link and export imports listed in the app. Administrators can stop new imports per provider and keep maintenance notes. Existing recipes stay available. Test public website checks basic availability only; provider blocks can make a probe fail even when a normal browser works. No private recipe account OAuth sync is configured; do not enter third-party passwords or invent callback URLs. Custom Home inventory sections can be created, renamed, reordered and archived. Archiving preserves existing items.')
 doc.add_heading('Reading the error catalog',1)
 doc.add_paragraph('Each row below gives an error code, the displayed message, and suggested checks. The paragraph above each group explains its common meaning. Braces such as {detail} represent a value supplied at runtime. The HTTP number describes the result class, not a guarantee that a retry is safe. Generic family codes cover failures that do not match a more specific message.')
 statuses=sorted(set(e['status'] for e in entries))
 def plain(text):return str(text).replace('’',"'").replace('–','-').replace('—','-').replace('…','...')
 for status in statuses:
  group=[e for e in entries if e['status']==status]
  title=admin_ops.FAMILIES[status][1] if status in admin_ops.FAMILIES else 'Browser error reports'
  doc.add_heading(title,1)
  if status in admin_ops.FAMILIES:doc.add_paragraph(plain(admin_ops.FAMILIES[status][2]))
  else:doc.add_paragraph('A browser script or asynchronous action failed. Server data may still be saved; refresh before retrying.')
  table=doc.add_table(rows=1,cols=3);table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
  for c,width,title in zip(table.rows[0].cells,[1.2,2.5,3.3],['Code','Displayed message','What to check']):c.width=Inches(width);c.text=title
  for e in group:
   cells=table.add_row().cells
   for cell,width,value in zip(cells,[1.2,2.5,3.3],[e['code'],e['message'],e['checks']]):
    cell.width=Inches(width);cell.text=plain(value)
  for index,row in enumerate(table.rows):
   for cell in row.cells:
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;props=cell._tc.get_or_add_tcPr()
    margins=OxmlElement('w:tcMar')
    for edge in ['top','left','bottom','right']:
     el=OxmlElement('w:'+edge);el.set(qn('w:w'),'90');el.set(qn('w:type'),'dxa');margins.append(el)
    props.append(margins);shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'DCE6F1' if index==0 else ('F5F7FA' if index%2 else 'FFFFFF'));props.append(shade)
    for para in cell.paragraphs:
     para.paragraph_format.space_after=Pt(3);para.paragraph_format.space_before=Pt(3)
     for run in para.runs:run.font.size=Pt(10);run.bold=index==0;run.font.color.rgb=RGBColor(0,0,0)
   trPr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trPr.append(cant)
  repeat=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(repeat)
  borders=OxmlElement('w:tblBorders')
  for edge in ['top','left','bottom','right','insideH','insideV']:
   el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
  table._tbl.tblPr.append(borders)
  doc.add_paragraph('')
 doc.core_properties.title='PantriPortal administrator error guide';doc.core_properties.author='PantriPortal';doc.core_properties.subject='Error codes activity logs and troubleshooting'
 dest=ROOT/'app/admin-error-guide.docx';doc.save(dest)
 print('Built',len(entries),'error references and Word guide:',dest)

if __name__=='__main__':document(build())
